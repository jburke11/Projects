import docx
import datetime
from docx.enum.text import WD_ALIGN_PARAGRAPH
document = docx.Document("/Users/joema/OneDrive/Documents/tc_labels.docx")
style = document.styles["Normal"]
font = style.font.bold=True
form = style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
table = document.tables[0]

s = input("Input all 's' lines").split(",") * 2
gwh = input ("Input all gwh lines").split(",") * 2
other = input ("input all other lines").split(",") * 2
date = datetime.datetime.now()
date = date.strftime("%x")
s = [i for i in s if i]
gwh = [i for i in gwh if i]
other = [i for i in other if i]
s.sort()
gwh.sort()
length = len(s) + len(gwh) + len(other)
rows = int(length / 4)
extra_cells = length % 4

for i in range(rows):
    table.add_row()
if extra_cells != 0:
    table.add_row()
row_num = 0
col_num = 0
for element in s:
    if col_num % 2 != 0:
        continue
    elif (col_num / 6 == 1):
        table.cell(row_num, col_num).text = "SUP19xM6" + "\t" + "# " + element + "\n" + "S" + "\t" + date
        row_num += 1
        col_num = 0
    else:
        table.cell(row_num, col_num).text = "SUP19xM6" + "\t" + "# " + element + "\n" + "S" + "\t" +  date
        col_num += 2
for element in gwh:
    if col_num % 2 != 0:
        continue
    elif (col_num / 6 == 1):
        table.cell(row_num, col_num).text = "SUP19xM6" + "\t" + "# " + element + "\n" + "GMH" + "\t" + date
        row_num += 1
        col_num = 0
    else:
        table.cell(row_num, col_num).text = "SUP19xM6" + "\t" + "# " + element + "\n" + "GMH" + "\t" + date
        col_num += 2
form = style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
for element in other:
    if col_num % 2 != 0:
        continue
    elif (col_num / 6 == 1):
        table.cell(row_num, col_num).text = element + "\t" + "\n" + date + "\t"
        row_num += 1
        col_num = 0
    else:
        table.cell(row_num, col_num).text = element + "\t" + "\n" + date + "\t"
        col_num += 2
document.save("test.docx")


